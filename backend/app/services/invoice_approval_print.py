from copy import copy
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


_APPROVAL_TEMPLATE = (
    Path(__file__).resolve().parent.parent / "templates" / "approval" / "invoice.docx"
)
_DETAIL_TEMPLATE = (
    Path(__file__).resolve().parent.parent / "templates" / "invoice" / "invoice_detail.xlsx"
)


def _replace_paragraph(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = text


def _date_text(value) -> str:
    return f"{value.year}年{value.month}月{value.day}日" if value else ""


def build_invoice_approval_docx(values: dict) -> bytes:
    document = Document(_APPROVAL_TEMPLATE)
    amount = Decimal(values.get("amount") or 0)
    replacements = {
        "购买方单位名称:": f"购买方单位名称: {values.get('customer_name', '')}",
        "纳税人识别号:": f"纳税人识别号: {values.get('tax_no', '')}",
        "地址、电话:": (
            f"地址、电话: {values.get('customer_address', '')} "
            f"{values.get('customer_phone', '')}"
        ).rstrip(),
        "开户行:": f"开户行: {values.get('customer_bank_name', '')}",
        "账号：": f"账号：{values.get('customer_bank_account', '')}",
        "开票金额:": (
            f"开票金额: {amount:.2f}    合同号: {values.get('contract_no', '')}"
        ),
        "开票事由:": f"开票事由: {values.get('business_item', '业务事项')}",
        "发票类型:": f"发票类型: {values.get('invoice_type', '')}",
        "申请人:": (
            f"申请人: {values.get('applicant_name', '')}    "
            f"部门负责人审批: {values.get('approver_name', '')} "
            f"{_date_text(values.get('approval_date'))}"
        ),
        "申请开票时间:": f"申请开票时间: {_date_text(values.get('apply_date'))}",
        "备注：": f"备注：{values.get('remark', '')}",
    }
    for paragraph in document.paragraphs:
        stripped = paragraph.text.strip()
        for prefix, replacement in replacements.items():
            if stripped.startswith(prefix):
                _replace_paragraph(paragraph, replacement)
                break
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _copy_row_style(worksheet, source_row: int, target_row: int) -> None:
    worksheet.row_dimensions[target_row].height = worksheet.row_dimensions[source_row].height
    for column in range(1, worksheet.max_column + 1):
        source = worksheet.cell(source_row, column)
        target = worksheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)


def build_invoice_detail_xlsx(details) -> bytes:
    workbook = load_workbook(_DETAIL_TEMPLATE)
    worksheet = workbook.active
    rows = list(details)
    detail_start = 2
    template_detail_count = 3
    total_row = detail_start + template_detail_count
    total_styles = [
        copy(worksheet.cell(total_row, column)._style)
        for column in range(1, 5)
    ]
    total_height = worksheet.row_dimensions[total_row].height

    if len(rows) > template_detail_count:
        worksheet.insert_rows(total_row, amount=len(rows) - template_detail_count)
        for row_number in range(
            detail_start + template_detail_count,
            detail_start + len(rows),
        ):
            _copy_row_style(worksheet, detail_start, row_number)
    elif len(rows) < template_detail_count:
        worksheet.delete_rows(
            detail_start + len(rows),
            amount=template_detail_count - len(rows),
        )

    total_row = detail_start + len(rows)
    worksheet.row_dimensions[total_row].height = total_height
    for column, style in enumerate(total_styles, start=1):
        worksheet.cell(total_row, column)._style = style

    for offset, detail in enumerate(rows):
        row_number = detail_start + offset
        worksheet.cell(row_number, 1, detail.get("item_name", ""))
        worksheet.cell(row_number, 2, detail.get("platform", ""))
        worksheet.cell(row_number, 3, "价税合计")
        worksheet.cell(row_number, 4, Decimal(detail.get("amount") or 0))

    worksheet.cell(total_row, 1, "合计")
    worksheet.cell(total_row, 2, None)
    worksheet.cell(total_row, 3, None)
    worksheet.cell(
        total_row,
        4,
        sum(
            (Decimal(detail.get("amount") or 0) for detail in rows),
            Decimal("0"),
        ),
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
