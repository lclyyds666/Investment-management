from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

import app.db.init_db  # noqa: F401
from app.api.deps import get_current_user
from app.core.enums import (
    AssignmentStatus,
    ContractStatus,
    ContractType,
    InvoiceApprovalStatus,
    InvoiceDirection,
    InvoiceSourceKind,
    InvoiceStatus,
    WorkflowTargetType,
    WorkflowTaskStatus,
)
from app.db.base import Base
from app.db.session import get_db
from app.main import create_app
from app.models.approval_form import ApprovalForm, ApprovalFormAction
from app.models.invoice import Invoice, InvoiceDetail
from app.models.organization import Organization, Position, UserAssignment
from app.models.user import User
from app.models.workflow import WorkflowInstance, WorkflowTask
from app.schemas.approval_form import ApprovalFormCreate, ApprovalFormUpdate
from app.services.organization_catalog import seed_authorization_catalog
from app.services.workflow_catalog import WORKFLOW_DEFINITIONS
from app.services.workflow_engine import seed_workflow_definitions


@pytest.fixture
def invoice_workflow_api():
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    event.listen(engine, "connect", lambda connection, _: connection.execute("PRAGMA foreign_keys=ON"))
    Base.metadata.create_all(engine)
    db = Session(engine)
    seed_authorization_catalog(db)

    def add_user(username, full_name, position_code):
        user = User(
            username=username,
            full_name=full_name,
            hashed_password="test",
            is_active=True,
        )
        db.add(user)
        db.flush()
        db.add(UserAssignment(
            user_id=user.id,
            organization_id=db.scalar(select(Organization.id).where(
                Organization.code == "supplymanagement"
            )),
            position_id=db.scalar(select(Position.id).where(Position.code == position_code)),
            valid_from=date(2026, 1, 1),
            status=AssignmentStatus.ACTIVE,
        ))
        db.flush()
        return user

    publisher = User(
        username="invoice-publisher",
        full_name="Publisher",
        hashed_password="test",
        is_active=True,
    )
    db.add(publisher)
    db.flush()
    users = {
        "reviewer": add_user("invoice-reviewer", "业务复核甲", "supply.business_reviewer"),
        "governance": add_user("invoice-governance", "供管负责人乙", "governance.supply_leader"),
        "handler": add_user("invoice-handler", "业务经办丙", "supply.business_handler"),
    }
    seed_workflow_definitions(db, publisher.id)
    invoice = Invoice(
        invoice_title="景区销项发票",
        tax_no="91370000123456789X",
        invoice_type="增值税专用发票",
        amount=Decimal("100.00"),
        status=InvoiceStatus.PENDING,
        direction=InvoiceDirection.OUTPUT,
        source_kind=InvoiceSourceKind.TICKET,
        scenic_id="scenic-a",
        period_key="2026-09",
        generated_by_source=True,
        customer_name="山东文旅客户有限公司",
        customer_social_credit_code="91370000123456789X",
        customer_address="济南市历下区经十路1号",
        customer_phone="0531-12345678",
        customer_bank_name="中国银行济南分行",
        customer_bank_account="1234567890",
        contract_no="HT-2026-001",
    )
    invoice.details.extend([
        InvoiceDetail(line_no=1, platform="携程", item_name="门票", amount=Decimal("40.00")),
        InvoiceDetail(line_no=2, platform="美团", item_name="门票", amount=Decimal("60.00")),
    ])
    db.add(invoice)
    db.commit()

    current = {"user": users["reviewer"]}
    app = create_app()
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: current["user"]
    client = TestClient(app, raise_server_exceptions=False)
    yield db, client, current, users, invoice
    client.close()
    app.dependency_overrides.clear()
    db.close()
    engine.dispose()


def test_invoice_catalog_has_exact_two_node_route():
    definition = next(item for item in WORKFLOW_DEFINITIONS if item.code == "supply.invoice.v1")

    assert definition.target_type == WorkflowTargetType.INVOICE_APPROVAL
    assert [(node.code, node.position_code) for node in definition.nodes] == [
        ("applicant", "supply.business_reviewer"),
        ("department_head", "governance.supply_leader"),
    ]
    assert definition.nodes[0].auto_complete_on_submit is True


def test_invoice_approval_two_node_permissions_and_projection(invoice_workflow_api):
    db, client, current, users, invoice = invoice_workflow_api

    created = client.post(f"/api/v1/invoices/{invoice.id}/approval-form")
    assert created.status_code == 200, created.text
    form_id = created.json()["data"]["id"]
    form = db.get(ApprovalForm, form_id)
    assert form.form_type == ContractType.INVOICE
    assert form.invoice_id == invoice.id
    assert form.created_by == users["reviewer"].id
    assert form.customer_name == invoice.customer_name
    assert form.contract_no == invoice.contract_no
    assert form.amount == invoice.amount

    submitted = client.post(f"/api/v1/invoices/{invoice.id}/approval-form/submit")
    assert submitted.status_code == 200, submitted.text
    db.refresh(form)
    instance = db.get(WorkflowInstance, form.workflow_instance_id)
    tasks = list(db.scalars(
        select(WorkflowTask)
        .where(WorkflowTask.instance_id == instance.id)
        .order_by(WorkflowTask.sequence)
    ))
    assert instance.target_type == WorkflowTargetType.INVOICE_APPROVAL
    assert [task.required_position_code for task in tasks] == [
        "supply.business_reviewer",
        "governance.supply_leader",
    ]
    assert tasks[0].status == WorkflowTaskStatus.APPROVED
    assert tasks[1].status == WorkflowTaskStatus.ACTIVE
    submit_projection = db.scalar(select(ApprovalFormAction).where(
        ApprovalFormAction.form_id == form.id
    ))
    assert submit_projection.approver_id == users["reviewer"].id
    assert submit_projection.position_code == "supply.business_reviewer"

    current["user"] = users["governance"]
    pending = client.get("/api/v1/approval/pending-count")
    assert pending.status_code == 200, pending.text
    assert pending.json()["data"]["business"] == 1

    current["user"] = users["handler"]
    denied = client.post(
        f"/api/v1/approval-forms/{form.id}/approve",
        json={"comment": "越权审批"},
    )
    assert denied.status_code == 403, denied.text

    current["user"] = users["governance"]
    approved = client.post(
        f"/api/v1/approval-forms/{form.id}/approve",
        json={"comment": "同意开票"},
    )
    assert approved.status_code == 200, approved.text
    db.refresh(form)
    db.refresh(invoice)
    assert form.status == ContractStatus.APPROVED
    assert invoice.approval_status == InvoiceApprovalStatus.APPROVED
    assert invoice.workflow_instance_id == form.workflow_instance_id
    printed = client.get(f"/api/v1/invoices/{invoice.id}/approval-form/print")
    assert printed.status_code == 200, printed.text
    document = Document(BytesIO(printed.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "业务复核甲" in text
    assert "供管负责人乙" in text
    invoice.details[0].amount = Decimal("30.00")
    db.commit()
    assert client.get(f"/api/v1/invoices/{invoice.id}/approval-form/print").status_code == 409


def test_rejected_invoice_form_reopens_and_resubmits_after_edits(invoice_workflow_api):
    db, client, current, users, invoice = invoice_workflow_api
    created = client.post(f"/api/v1/invoices/{invoice.id}/approval-form")
    form_id = created.json()["data"]["id"]
    submitted = client.post(f"/api/v1/invoices/{invoice.id}/approval-form/submit")
    instance_id = submitted.json()["data"]["workflow_instance_id"]

    current["user"] = users["governance"]
    rejected = client.post(
        f"/api/v1/approval-forms/{form_id}/reject",
        json={"comment": "请更新合同号"},
    )
    assert rejected.status_code == 200, rejected.text

    current["user"] = users["reviewer"]
    updated = client.put(
        f"/api/v1/invoices/{invoice.id}",
        json={"contract_no": "HT-2026-002"},
    )
    assert updated.status_code == 200, updated.text
    reopened = client.post(f"/api/v1/invoices/{invoice.id}/approval-form")
    assert reopened.status_code == 200, reopened.text
    assert reopened.json()["data"]["id"] == form_id
    assert reopened.json()["data"]["contract_no"] == "HT-2026-002"
    resumed = client.post(f"/api/v1/invoices/{invoice.id}/approval-form/submit")
    assert resumed.status_code == 200, resumed.text
    assert resumed.json()["data"]["workflow_instance_id"] == instance_id
    assert resumed.json()["data"]["status"] == ContractStatus.PENDING.value
    db.refresh(invoice)
    assert invoice.approval_status == InvoiceApprovalStatus.PENDING


def test_invoice_approval_creation_requires_manage_data_and_balanced_details(invoice_workflow_api):
    db, client, current, users, invoice = invoice_workflow_api
    current["user"] = users["handler"]
    assert client.post(f"/api/v1/invoices/{invoice.id}/approval-form").status_code == 403

    current["user"] = users["reviewer"]
    invoice.customer_social_credit_code = ""
    db.commit()
    missing_tax = client.post(f"/api/v1/invoices/{invoice.id}/approval-form")
    assert missing_tax.status_code == 422, missing_tax.text

    invoice.customer_social_credit_code = "91370000123456789X"
    invoice.details[0].amount = Decimal("30.00")
    db.commit()
    unbalanced = client.post(f"/api/v1/invoices/{invoice.id}/approval-form")
    assert unbalanced.status_code == 409, unbalanced.text


def test_invoice_forms_cannot_use_generic_crud_or_submit(invoice_workflow_api):
    db, client, current, users, invoice = invoice_workflow_api
    assert "invoice_id" not in ApprovalFormCreate.model_fields
    assert "invoice_id" not in ApprovalFormUpdate.model_fields
    current["user"] = users["handler"]
    bypass = client.post("/api/v1/approval-forms", json={
        "form_type": "invoice",
        "invoice_id": invoice.id,
        "customer_name": "绕过校验的客户",
        "amount": "1.00",
    })
    assert bypass.status_code == 403, bypass.text
    assert db.scalar(select(ApprovalForm).where(
        ApprovalForm.invoice_id == invoice.id
    )) is None

    current["user"] = users["reviewer"]
    first = client.post(f"/api/v1/invoices/{invoice.id}/approval-form")
    second = client.post(f"/api/v1/invoices/{invoice.id}/approval-form")
    assert first.status_code == 200, first.text
    assert second.status_code == 200, second.text
    form_id = first.json()["data"]["id"]
    assert second.json()["data"]["id"] == form_id
    assert len(list(db.scalars(select(ApprovalForm).where(
        ApprovalForm.invoice_id == invoice.id
    )))) == 1

    current["user"] = users["handler"]
    assert client.put(
        f"/api/v1/approval-forms/{form_id}",
        json={"customer_name": "绕过修改"},
    ).status_code == 403
    assert client.post(
        f"/api/v1/approval-forms/{form_id}/submit",
        json={},
    ).status_code == 403
    assert client.delete(f"/api/v1/approval-forms/{form_id}").status_code == 403
    assert db.get(ApprovalForm, form_id) is not None


def test_approved_invoice_docx_uses_immutable_snapshots(invoice_workflow_api):
    db, client, current, users, invoice = invoice_workflow_api
    created = client.post(f"/api/v1/invoices/{invoice.id}/approval-form")
    form_id = created.json()["data"]["id"]
    assert client.post(
        f"/api/v1/invoices/{invoice.id}/approval-form/submit"
    ).status_code == 200
    current["user"] = users["governance"]
    assert client.post(
        f"/api/v1/approval-forms/{form_id}/approve",
        json={"comment": "同意"},
    ).status_code == 200

    users["reviewer"].full_name = "已改名申请人"
    users["governance"].full_name = "已改名负责人"
    invoice.customer_name = "变更后客户"
    invoice.customer_social_credit_code = "NEW-TAX-NO"
    invoice.customer_address = "变更后地址"
    invoice.customer_phone = "999999"
    invoice.customer_bank_name = "变更后银行"
    invoice.customer_bank_account = "000000"
    invoice.amount = Decimal("999.00")
    invoice.contract_no = "NEW-CONTRACT"
    invoice.invoice_type = "增值税普通发票"
    invoice.details[0].amount = Decimal("400.00")
    invoice.details[1].amount = Decimal("599.00")
    db.commit()

    printed = client.get(f"/api/v1/invoices/{invoice.id}/approval-form/print")
    assert printed.status_code == 200, printed.text
    document = Document(BytesIO(printed.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for original in (
        "业务复核甲",
        "供管负责人乙",
        "山东文旅客户有限公司",
        "91370000123456789X",
        "济南市历下区经十路1号",
        "0531-12345678",
        "中国银行济南分行",
        "1234567890",
        "100.00",
        "HT-2026-001",
        "增值税专用发票",
    ):
        assert original in text
    for changed in (
        "已改名申请人",
        "已改名负责人",
        "变更后客户",
        "NEW-TAX-NO",
        "变更后地址",
        "999999",
        "变更后银行",
        "000000",
        "999.00",
        "NEW-CONTRACT",
        "增值税普通发票",
    ):
        assert changed not in text


def test_invoice_print_endpoints_enforce_approval_and_balance(invoice_workflow_api):
    db, client, current, users, invoice = invoice_workflow_api
    created = client.post(f"/api/v1/invoices/{invoice.id}/approval-form")
    assert created.status_code == 200, created.text
    assert client.get(f"/api/v1/invoices/{invoice.id}/approval-form/print").status_code == 409

    detail_print = client.get(f"/api/v1/invoices/{invoice.id}/details/print")
    assert detail_print.status_code == 200, detail_print.text
    assert detail_print.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    invoice.details[0].amount = Decimal("30.00")
    db.commit()
    assert client.get(f"/api/v1/invoices/{invoice.id}/details/print").status_code == 409
