from datetime import date
from decimal import Decimal
from io import BytesIO

from docx import Document
from openpyxl import load_workbook

from app.services.invoice_approval_print import (
    build_invoice_approval_docx,
    build_invoice_detail_xlsx,
)


def test_invoice_approval_docx_replaces_all_business_fields():
    data = build_invoice_approval_docx({
        "customer_name": "山东文旅客户有限公司",
        "tax_no": "91370000123456789X",
        "customer_address": "济南市历下区经十路1号",
        "customer_phone": "0531-12345678",
        "customer_bank_name": "中国银行济南分行",
        "customer_bank_account": "1234567890",
        "amount": Decimal("100.00"),
        "contract_no": "HT-2026-001",
        "business_item": "业务事项",
        "invoice_type": "增值税专用发票",
        "applicant_name": "业务复核甲",
        "approver_name": "供管负责人乙",
        "apply_date": date(2026, 9, 28),
        "approval_date": date(2026, 9, 29),
        "remark": "测试备注",
    })

    document = Document(BytesIO(data))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for value in (
        "山东文旅客户有限公司",
        "91370000123456789X",
        "济南市历下区经十路1号",
        "0531-12345678",
        "中国银行济南分行",
        "1234567890",
        "100.00",
        "HT-2026-001",
        "业务事项",
        "业务复核甲",
        "供管负责人乙",
        "2026年9月28日",
        "2026年9月29日",
    ):
        assert value in text


def test_invoice_detail_xlsx_inserts_rows_and_writes_total():
    details = [
        {"item_name": "门票", "platform": "携程", "amount": Decimal("10.00")},
        {"item_name": "门票", "platform": "美团", "amount": Decimal("20.00")},
        {"item_name": "酒店", "platform": "抖音", "amount": Decimal("30.00")},
        {"item_name": "酒店", "platform": "票付通", "amount": Decimal("40.00")},
    ]

    data = build_invoice_detail_xlsx(details)

    workbook = load_workbook(BytesIO(data), data_only=True)
    worksheet = workbook.active
    assert worksheet.max_row == 6
    assert [worksheet.cell(row=row, column=2).value for row in range(2, 6)] == [
        "携程", "美团", "抖音", "票付通",
    ]
    assert [worksheet.cell(row=row, column=4).value for row in range(2, 6)] == [
        10, 20, 30, 40,
    ]
    assert worksheet["A6"].value == "合计"
    assert worksheet["D6"].value == 100
    assert [str(item) for item in worksheet.merged_cells.ranges] == ["A1:D1"]
